"""
lab_plotter.py
"""
import numpy as np
from openpyxl import load_workbook
import csv


def read_excel(path, sheet=0, columns=None, skip_rows=0):
    """
    Read data from an Excel file into a dict of numpy arrays.

    Args:
        path:      path to .xlsx file
        sheet:     sheet name (str) or index (int)
        columns:   dict mapping your names to column letters
                   e.g. {'freq': 'A', 'gain': 'B', 'phase': 'C'}
                   if None, reads all columns using row 1 as headers
        skip_rows: number of rows to skip before data starts (after header)

    Returns:
        dict of numpy arrays keyed by your column names

    Example:
        data = read_excel('lab06.xlsx', sheet='Sheet1',
                          columns={'freq': 'A', 'vo': 'B', 'vi': 'C', 'phase': 'F'},
                          skip_rows=1)
        print(data['freq'])
        print(data['vo'])
    """
    wb = load_workbook(path, data_only=True)

    if isinstance(sheet, int):
        ws = wb.worksheets[sheet]
    else:
        if sheet not in wb.sheetnames:
            print(f"\n  Sheet '{sheet}' not found.")
            print(f"  Available sheets: {wb.sheetnames}")
            wb.close()
            return {}
        ws = wb[sheet]

    if columns is None:
        columns = {}
        for cell in ws[1]:
            if cell.value is not None:
                columns[str(cell.value)] = cell.column_letter

    result = {}
    for name, col_letter in columns.items():
        values = []
        col_idx = ws[col_letter + '1'].column
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx,
                                min_row=2 + skip_rows):
            cell = row[0]
            if cell.value is not None:
                try:
                    values.append(float(cell.value))
                except (ValueError, TypeError):
                    values.append(np.nan)
        result[name] = np.array(values)

    wb.close()

    # drop rows where any column is NaN
    if result:
        mask = np.ones(len(next(iter(result.values()))), dtype=bool)
        for arr in result.values():
            mask &= ~np.isnan(arr)
        result = {k: v[mask] for k, v in result.items()}

    return result


def read_csv_multisim(path):
    """
    Read a Multisim AC Sweep CSV export.
    Auto-detects single-probe (3 col) or dual-probe (7 col) format.

    Returns:
        dict with 'freq', 'mag', 'phase' as numpy arrays
        (mag is raw ratio, not dB)
    """
    freq, mag, phase = [], [], []

    with open(path) as f:
        reader = csv.reader(f)
        header = next(reader)
        n = len(header)

        for row in reader:
            if not row or not row[0].strip():
                continue
            try:
                if n >= 7:
                    freq.append(float(row[4]))
                    mag.append(float(row[5]))
                    phase.append(float(row[6]))
                else:
                    freq.append(float(row[0]))
                    mag.append(float(row[1]))
                    phase.append(float(row[2]))
            except (ValueError, IndexError):
                continue

    return {
        'freq': np.array(freq),
        'mag': np.array(mag),
        'phase': np.array(phase)
    }


# ─────────────────────────────────────────
# CONVERSIONS
# ─────────────────────────────────────────

def to_db(ratio):
    """Voltage ratio → dB. Works on numpy arrays or floats."""
    return 20 * np.log10(np.where(np.asarray(ratio) > 0, ratio, 1e-12))


def cutoff_freq(R, C):
    """Returns 1/(2πRC) in Hz."""
    return 1 / (2 * np.pi * R * C)


# ─────────────────────────────────────────
# THEORY CURVES
# ─────────────────────────────────────────

def rc_lowpass(freq, R, C):
    """
    Passive RC lowpass transfer function.
    Returns dict with 'freq', 'gain_db', 'phase'.
    """
    f = np.asarray(freq)
    wRC = 2 * np.pi * f * R * C
    gain_db = to_db(1 / np.sqrt(1 + wRC**2))
    phase = -np.degrees(np.arctan(wRC))
    return {'freq': f, 'gain_db': gain_db, 'phase': phase}


def rc_highpass(freq, R, C):
    """Passive RC highpass transfer function."""
    f = np.asarray(freq)
    wRC = 2 * np.pi * f * R * C
    gain_db = to_db(wRC / np.sqrt(1 + wRC**2))
    phase = np.degrees(np.arctan(1 / wRC))
    return {'freq': f, 'gain_db': gain_db, 'phase': phase}


def inv_lowpass(freq, R1, R2, C2):
    """
    Inverting op-amp lowpass: H = -(R2/R1) / (1 + jwR2C2)
    Phase includes the 180° from inversion.
    """
    f = np.asarray(freq)
    G = R2 / R1
    wRC = 2 * np.pi * f * R2 * C2
    gain_db = to_db(G / np.sqrt(1 + wRC**2))
    phase = 180 - np.degrees(np.arctan(wRC))
    return {'freq': f, 'gain_db': gain_db, 'phase': phase}


def logspace_freq(f_start, f_stop, n_per_decade=50):
    """Generate log-spaced frequency points for smooth theory curves."""
    return np.logspace(np.log10(f_start), np.log10(f_stop),
                       int(n_per_decade * np.log10(f_stop / f_start)))


# ─────────────────────────────────────────
# PLOTTING
# ─────────────────────────────────────────

import matplotlib.pyplot as plt


def bode(datasets, title='Bode Plot', save=None,
         gain_range=None, phase_range=None, freq_range=None):
    """
    Bode plot with any number of overlaid datasets.

    Args:
        datasets: list of dicts, each with:
            'freq':    array of frequencies in Hz
            'gain_db': array of gain in dB
            'phase':   array of phase in degrees (can have None/NaN)
            'label':   string for legend
            'style':   'sim' → red solid | 'meas' → blue dashed + markers
                       'theory' → black solid  (optional, defaults to 'sim')
        title: plot title
        save:  filename to save PNG (None = just show)

    Example:
        sim = read_csv_multisim('sim.csv')
        sim_data = {
            'freq': sim['freq'],
            'gain_db': to_db(sim['mag']),
            'phase': sim['phase'],
            'label': 'Simulation',
            'style': 'sim'
        }
        bode([sim_data, meas_data], title='Lab 06 Q2', save='q2.png')
    """
    styles = {
        'sim':     {'color': 'red',   'ls': '-',  'lw': 2,   'marker': None, 'ms': 0},
        'meas':    {'color': 'blue',  'ls': '--', 'lw': 1.5, 'marker': 'o',  'ms': 5},
        'theory':  {'color': 'black', 'ls': '-',  'lw': 1.5, 'marker': None, 'ms': 0},
    }

    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(10, 7), sharex=True)
    fig.suptitle(title, fontsize=13, fontweight='bold')

    for ds in datasets:
        s = styles.get(ds.get('style', 'sim'), styles['sim'])

        # gain
        ax1.semilogx(ds['freq'], ds['gain_db'], label=ds.get('label', ''),
                     color=s['color'], ls=s['ls'], lw=s['lw'],
                     marker=s['marker'], ms=s['ms'])

        # phase — filter NaN/None
        p = np.asarray(ds['phase'], dtype=float)
        valid = ~np.isnan(p)
        if np.any(valid):
            ax2.semilogx(np.asarray(ds['freq'])[valid], p[valid],
                         label=ds.get('label', ''),
                         color=s['color'], ls=s['ls'], lw=s['lw'],
                         marker=s['marker'], ms=s['ms'])

    ax1.set_ylabel('Gain (dB)')
    ax1.legend()
    ax1.grid(True, which='both', alpha=0.3)
    if gain_range: ax1.set_ylim(gain_range)

    ax2.set_xlabel('Frequency (Hz)')
    ax2.set_ylabel('Phase (°)')
    ax2.legend()
    ax2.grid(True, which='both', alpha=0.3)
    if phase_range: ax2.set_ylim(phase_range)
    if freq_range:  ax2.set_xlim(freq_range)

    plt.tight_layout()
    if save:
        fig.savefig(save, dpi=150, bbox_inches='tight')
        print(f'Saved: {save}')
    plt.show()


# ─────────────────────────────────────────
# WRITER — dump results back to Excel
# ─────────────────────────────────────────

from openpyxl import Workbook


def write_excel(path, sheets):
    """
    Write data to an Excel workbook.

    Args:
        path:   output .xlsx path
        sheets: dict mapping sheet names to dicts of arrays
                e.g. {'Simulation': {'freq': [...], 'gain_db': [...], 'phase': [...]},
                       'Measurements': {'freq': [...], 'vo': [...], 'vi': [...]}}

    Example:
        theory = rc_lowpass(logspace_freq(100, 1e6), 1800, 10e-9)
        write_excel('lab06_data.xlsx', {
            'Theory': theory,
            'Measurements': {'freq': meas_f, 'vo': meas_vo, 'vi': meas_vi}
        })
    """
    wb = Workbook()
    first = True

    for sheet_name, data in sheets.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(sheet_name)

        # headers
        headers = list(data.keys())
        for c, h in enumerate(headers, 1):
            ws.cell(row=1, column=c, value=h)

        # data rows
        n_rows = len(next(iter(data.values())))
        for r in range(n_rows):
            for c, h in enumerate(headers, 1):
                val = data[h][r]
                ws.cell(row=r + 2, column=c, value=float(val) if not np.isnan(val) else None)

    wb.save(path)
    print(f'Saved: {path}')


if __name__ == '__main__':

    print("── Lab Plotter ──")
    sim_path = input("Multisim CSV path: ").strip()
    excel_path = input("Excel file path: ").strip()

    # show available sheets
    wb_preview = load_workbook(excel_path, data_only=True)
    print("\nAvailable sheets:")
    for i, name in enumerate(wb_preview.sheetnames):
        print(f"  [{i}] {name}")

    # pick gain sheet and show its columns
    gain_idx = int(input("\nSheet number for gain data: ").strip())
    ws_preview = wb_preview.worksheets[gain_idx]
    print("Columns:")
    for cell in ws_preview[1]:
        if cell.value is not None:
            print(f"  {cell.column_letter}: {cell.value}")

    col_freq = input("Column letter for frequency: ").strip()
    col_gain = input("Column letter for gain (dB): ").strip()

    # pick phase sheet
    phase_idx = int(input("\nSheet number for phase data (same number if same sheet): ").strip())
    if phase_idx != gain_idx:
        ws_phase = wb_preview.worksheets[phase_idx]
        print("Columns:")
        for cell in ws_phase[1]:
            if cell.value is not None:
                print(f"  {cell.column_letter}: {cell.value}")
    col_phase = input("Column letter for phase: ").strip()
    wb_preview.close()

    save_name = input("\nSave plot as (e.g. lab06_Q2.png): ").strip()
    title = input("Plot title: ").strip()

    # load
    sim = read_csv_multisim(sim_path)
    meas = read_excel(excel_path, sheet=gain_idx,
                      columns={'freq': col_freq, 'gain_db': col_gain})
    meas_phase = read_excel(excel_path, sheet=phase_idx,
                            columns={'freq': col_freq, 'phase': col_phase})

    # plot
    bode([
        {'freq': sim['freq'], 'gain_db': to_db(sim['mag']), 'phase': sim['phase'],
         'label': 'Simulation', 'style': 'sim'},
        {'freq': meas['freq'], 'gain_db': meas['gain_db'], 'phase': meas_phase['phase'],
         'label': 'Measurement', 'style': 'meas'},
    ], title=title, save=save_name)